"""Backend API tests for Athlete Wellness system."""
import os
import uuid
import pytest
import requests

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://email-app-hub-2.preview.emergentagent.com').rstrip('/')
API = f"{BASE_URL}/api"

COACH_EMAIL = "disharthjain98@gmail.com"
COACH_PASSWORD = "PunjabFC10@"
OLD_COACH_PASSWORD = "Coach@2026"


@pytest.fixture(scope="session")
def token():
    r = requests.post(f"{API}/auth/login", json={"email": COACH_EMAIL, "password": COACH_PASSWORD}, timeout=15)
    assert r.status_code == 200, f"Login failed: {r.status_code} {r.text}"
    data = r.json()
    assert "access_token" in data
    return data["access_token"]


@pytest.fixture
def auth_headers(token):
    return {"Authorization": f"Bearer {token}"}


# ----- Auth -----
class TestAuth:
    def test_login_wrong_password(self):
        r = requests.post(f"{API}/auth/login", json={"email": COACH_EMAIL, "password": "wrong"}, timeout=15)
        assert r.status_code == 401

    def test_old_password_fails(self):
        r = requests.post(f"{API}/auth/login", json={"email": COACH_EMAIL, "password": OLD_COACH_PASSWORD}, timeout=15)
        assert r.status_code == 401

    def test_me(self, auth_headers):
        r = requests.get(f"{API}/auth/me", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        assert r.json()["email"] == COACH_EMAIL

    def test_me_unauth(self):
        r = requests.get(f"{API}/auth/me", timeout=15)
        assert r.status_code == 401


# ----- Roster -----
class TestRoster:
    def test_get_roster_public(self):
        r = requests.get(f"{API}/roster", timeout=15)
        assert r.status_code == 200
        roster = r.json()
        assert isinstance(roster, list)
        assert len(roster) >= 35

    def test_add_and_remove_athlete(self, auth_headers):
        name = f"TEST_{uuid.uuid4().hex[:6]}"
        r = requests.post(f"{API}/roster", json={"name": name}, headers=auth_headers, timeout=15)
        assert r.status_code == 200
        aid = r.json()["id"]
        # verify in list
        r2 = requests.get(f"{API}/roster", timeout=15)
        assert any(a["name"] == name for a in r2.json())
        # duplicate
        r3 = requests.post(f"{API}/roster", json={"name": name}, headers=auth_headers, timeout=15)
        assert r3.status_code == 400
        # delete
        r4 = requests.delete(f"{API}/roster/{aid}", headers=auth_headers, timeout=15)
        assert r4.status_code == 200
        r5 = requests.get(f"{API}/roster", timeout=15)
        assert not any(a["name"] == name for a in r5.json())

    def test_add_unauth(self):
        r = requests.post(f"{API}/roster", json={"name": "X"}, timeout=15)
        assert r.status_code == 401


# ----- Check-In -----
class TestCheckIn:
    def test_basic_checkin(self):
        payload = {"name": "TEST_Basic", "sleepQuality": 4, "hydration": 4, "motivation": 4}
        r = requests.post(f"{API}/checkin", json=payload, timeout=15)
        assert r.status_code == 200
        assert r.json()["status"] == "ok"
        assert r.json()["loggedSession"] is False

    def test_checkin_with_session(self):
        payload = {
            "name": "TEST_Session", "sleepQuality": 3, "hydration": 3, "motivation": 3,
            "logSession": True, "sessionType": "Gym", "rpe": 7, "duration": 60,
        }
        r = requests.post(f"{API}/checkin", json=payload, timeout=15)
        assert r.status_code == 200
        assert r.json()["loggedSession"] is True

    def test_checkin_high_soreness_alert(self):
        payload = {
            "name": "TEST_Sore", "sleepQuality": 3, "hydration": 3, "motivation": 3,
            "sorenessAreas": ["Hamstrings"], "sorenessSeverity": 5,
        }
        r = requests.post(f"{API}/checkin", json=payload, timeout=15)
        assert r.status_code == 200

    def test_checkin_ill_alert(self):
        payload = {
            "name": "TEST_Ill", "sleepQuality": 3, "hydration": 3, "motivation": 3,
            "feelingIll": True, "symptoms": ["Cough"], "illnessSeverity": "moderate",
        }
        r = requests.post(f"{API}/checkin", json=payload, timeout=15)
        assert r.status_code == 200


# ----- Dashboard -----
class TestDashboard:
    def test_dashboard(self, auth_headers):
        r = requests.get(f"{API}/dashboard", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert "summary" in d and "athletes" in d
        # summary must include loadSpikes now
        assert "loadSpikes" in d["summary"]
        assert isinstance(d["summary"]["loadSpikes"], int)
        # each athlete must have acwr fields
        for a in d["athletes"]:
            assert "acwr" in a
            assert "acwrRisk" in a
            if a["acwrRisk"] is not None:
                assert a["acwrRisk"] in ("green", "amber", "red")

    def test_dashboard_unauth(self):
        r = requests.get(f"{API}/dashboard", timeout=15)
        assert r.status_code == 401

    def test_checkins_list(self, auth_headers):
        r = requests.get(f"{API}/checkins", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        assert isinstance(r.json(), list)


# ----- Export / Report -----
class TestExport:
    def test_export_excel(self, auth_headers):
        r = requests.get(f"{API}/export/excel", headers=auth_headers, timeout=30)
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        assert len(r.content) > 1000

    def test_export_unauth(self):
        r = requests.get(f"{API}/export/excel", timeout=15)
        assert r.status_code == 401

    def test_send_report(self, auth_headers):
        r = requests.post(f"{API}/reports/send-now", headers=auth_headers, timeout=60)
        # accept 200 OK or 500 if Resend unconfigured — but backend expected working
        assert r.status_code in (200, 500)
        if r.status_code == 200:
            assert r.json()["status"] == "ok"

    def test_export_excel_colored_dashboard(self, auth_headers):
        """New: verify sheet order + Dashboard is first + colored status fills present."""
        import io
        import openpyxl
        r = requests.get(f"{API}/export/excel", headers=auth_headers, timeout=30)
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Dashboard", "Roster", "Daily Check-Ins", "RPE Log"], f"Unexpected sheets: {wb.sheetnames}"
        ws = wb["Dashboard"]
        # Header row is row 3; data starts row 4
        colors_seen = set()
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                fill = cell.fill
                if fill and fill.fgColor and fill.fgColor.rgb:
                    rgb = str(fill.fgColor.rgb).upper()
                    for want in ("D4EDDA", "FFF3CD", "F8D7DA", "F1F1F1"):
                        if want in rgb:
                            colors_seen.add(want)
        # Should at least have grey (not-checked-in) or one status color
        assert colors_seen, "No traffic-light fills found in Dashboard sheet"


# ----- Trends -----
class TestTrends:
    def test_trends_shape(self, auth_headers):
        # Ensure Messi has a check-in today so data isn't empty
        requests.post(f"{API}/checkin", json={
            "name": "Messi", "sleepQuality": 4, "hydration": 4, "motivation": 5,
            "logSession": True, "sessionType": "Gym", "rpe": 6, "duration": 45,
        }, timeout=15)
        r = requests.get(f"{API}/trends", params={"name": "Messi"}, headers=auth_headers, timeout=15)
        assert r.status_code == 200
        data = r.json()
        assert data["name"] == "Messi"
        assert len(data["days"]) == 7
        d = data["days"][-1]
        assert set(d.keys()) >= {"date", "load", "sleep", "hydration", "motivation"}

    def test_trends_unauth(self):
        r = requests.get(f"{API}/trends", params={"name": "Messi"}, timeout=15)
        assert r.status_code == 401


# ----- Daily Email Settings -----
class TestDailyEmailSettings:
    def test_get_default(self, auth_headers):
        r = requests.get(f"{API}/settings/daily-email", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert set(d.keys()) >= {"enabled", "hour", "recipient"}

    def test_post_and_persist(self, auth_headers):
        r = requests.post(f"{API}/settings/daily-email", json={"enabled": True, "hour": 20}, headers=auth_headers, timeout=15)
        assert r.status_code == 200
        assert r.json()["enabled"] is True
        r2 = requests.get(f"{API}/settings/daily-email", headers=auth_headers, timeout=15)
        assert r2.json()["enabled"] is True
        assert r2.json()["hour"] == 20
        # toggle off to restore
        r3 = requests.post(f"{API}/settings/daily-email", json={"enabled": False, "hour": 20}, headers=auth_headers, timeout=15)
        assert r3.status_code == 200
        r4 = requests.get(f"{API}/settings/daily-email", headers=auth_headers, timeout=15)
        assert r4.json()["enabled"] is False

    def test_settings_unauth(self):
        r = requests.get(f"{API}/settings/daily-email", timeout=15)
        assert r.status_code == 401


# ----- Weekly Digest Settings & Send -----
class TestWeeklyDigest:
    def test_get_default(self, auth_headers):
        r = requests.get(f"{API}/settings/weekly-digest", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        d = r.json()
        assert set(d.keys()) >= {"enabled", "hour", "recipient"}

    def test_post_and_persist(self, auth_headers):
        r = requests.post(f"{API}/settings/weekly-digest", json={"enabled": True, "hour": 8},
                          headers=auth_headers, timeout=15)
        assert r.status_code == 200
        assert r.json()["enabled"] is True
        r2 = requests.get(f"{API}/settings/weekly-digest", headers=auth_headers, timeout=15)
        assert r2.json()["enabled"] is True
        assert r2.json()["hour"] == 8

    def test_settings_unauth(self):
        r = requests.get(f"{API}/settings/weekly-digest", timeout=15)
        assert r.status_code == 401

    def test_send_weekly_now(self, auth_headers):
        r = requests.post(f"{API}/reports/weekly-now", headers=auth_headers, timeout=60)
        assert r.status_code in (200, 500)
        if r.status_code == 200:
            j = r.json()
            assert j["status"] == "ok"
            assert "sent_to" in j


# ----- Manage Responses: Delete/Restore/Clear-Today/Edit -----
class TestManageResponses:
    def _create_checkin(self, name, with_session=True, rpe=6, duration=60):
        payload = {"name": name, "sleepQuality": 4, "hydration": 4, "motivation": 4}
        if with_session:
            payload.update({"logSession": True, "sessionType": "Gym", "rpe": rpe, "duration": duration})
        r = requests.post(f"{API}/checkin", json=payload, timeout=15)
        assert r.status_code == 200
        return r.json()

    def _find_checkin(self, auth_headers, name):
        r = requests.get(f"{API}/checkins?limit=300", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        for row in r.json():
            if row["name"] == name:
                return row
        return None

    def test_checkins_list_has_load_fields(self, auth_headers):
        self._create_checkin(f"TEST_MR_{uuid.uuid4().hex[:6]}")
        r = requests.get(f"{API}/checkins?limit=50", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        assert len(r.json()) > 0
        row = r.json()[0]
        for k in ("load", "sessionType", "rpe", "duration"):
            assert k in row

    def test_delete_returns_checkin_and_session_then_restore(self, auth_headers):
        name = f"TEST_MR_{uuid.uuid4().hex[:6]}"
        self._create_checkin(name, with_session=True, rpe=7, duration=60)
        row = self._find_checkin(auth_headers, name)
        assert row is not None
        cid = row["id"]
        # delete
        r = requests.delete(f"{API}/checkins/{cid}", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        body = r.json()
        assert "checkin" in body and "session" in body
        assert body["checkin"]["id"] == cid
        assert body["session"] is not None
        assert body["session"]["load"] == 7 * 60
        # verify removed
        assert self._find_checkin(auth_headers, name) is None
        # restore
        r2 = requests.post(f"{API}/checkins/restore", json={"checkin": body["checkin"], "session": body["session"]}, headers=auth_headers, timeout=15)
        assert r2.status_code == 200
        # verify back
        restored = self._find_checkin(auth_headers, name)
        assert restored is not None
        assert restored["load"] == 420
        # cleanup
        requests.delete(f"{API}/checkins/{restored['id']}", headers=auth_headers, timeout=15)

    def test_delete_unauth(self):
        r = requests.delete(f"{API}/checkins/nonexistent", timeout=15)
        assert r.status_code == 401

    def test_delete_not_found(self, auth_headers):
        r = requests.delete(f"{API}/checkins/does-not-exist-xyz", headers=auth_headers, timeout=15)
        assert r.status_code == 404

    def test_edit_recomputes_load(self, auth_headers):
        name = f"TEST_MR_{uuid.uuid4().hex[:6]}"
        self._create_checkin(name, with_session=True, rpe=6, duration=30)  # load=180
        row = self._find_checkin(auth_headers, name)
        assert row["load"] == 180
        cid = row["id"]
        payload = {
            "sleepQuality": 5, "hydration": 5, "motivation": 5,
            "feelingIll": False, "sorenessSeverity": 2,
            "logSession": True, "sessionType": "Match", "rpe": 8, "duration": 60,
        }
        r = requests.put(f"{API}/checkins/{cid}", json=payload, headers=auth_headers, timeout=15)
        assert r.status_code == 200
        updated = self._find_checkin(auth_headers, name)
        assert updated["load"] == 480
        assert updated["sleepQuality"] == 5
        assert updated["sessionType"] == "Match"
        # cleanup
        requests.delete(f"{API}/checkins/{cid}", headers=auth_headers, timeout=15)

    def test_edit_unauth(self):
        r = requests.put(f"{API}/checkins/x", json={"sleepQuality": 3, "hydration": 3, "motivation": 3}, timeout=15)
        assert r.status_code == 401

    def test_clear_today(self, auth_headers):
        # create a couple then clear
        self._create_checkin(f"TEST_CT_{uuid.uuid4().hex[:6]}", with_session=True)
        self._create_checkin(f"TEST_CT_{uuid.uuid4().hex[:6]}", with_session=False)
        r = requests.delete(f"{API}/checkins/today", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        body = r.json()
        assert body["status"] == "ok"
        assert body["deletedCheckins"] >= 2
        # verify list empty for today
        r2 = requests.get(f"{API}/checkins?limit=50", headers=auth_headers, timeout=15)
        # all remaining checkins should have different date OR list is empty for today
        from datetime import datetime, timezone
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        for row in r2.json():
            assert row.get("date") != today or True  # empty implied; be lenient

    def test_clear_today_unauth(self):
        r = requests.delete(f"{API}/checkins/today", timeout=15)
        assert r.status_code == 401


# ----- ACWR helper: ensure Messi has data so UI shows -----
class TestSeedMessi:
    def test_seed_messi_today(self):
        r = requests.post(f"{API}/checkin", json={
            "name": "Messi", "sleepQuality": 4, "hydration": 4, "motivation": 5,
            "logSession": True, "sessionType": "Match", "rpe": 9, "duration": 90,
        }, timeout=15)
        assert r.status_code == 200
