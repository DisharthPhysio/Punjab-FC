from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import io
import re
import statistics
import uuid
import logging
import asyncio
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo
from typing import List, Optional, Literal

import jwt
import bcrypt
import resend
import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from fastapi import FastAPI, APIRouter, HTTPException, Request, Depends
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from apscheduler.schedulers.asyncio import AsyncIOScheduler

# ---------- Config ----------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
if not JWT_SECRET:
    raise RuntimeError("JWT_SECRET must be set to a long random string")
JWT_ALGORITHM = "HS256"
COACH_EMAIL = os.environ.get('COACH_EMAIL', '')
# Country code assumed for phone numbers typed as 10 digits (WhatsApp reminders). India = 91.
DEFAULT_COUNTRY_CODE = os.environ.get('DEFAULT_COUNTRY_CODE', '91').strip().lstrip('+') or '91'
SENDER_EMAIL = os.environ.get('SENDER_EMAIL', 'onboarding@resend.dev')
resend.api_key = os.environ.get('RESEND_API_KEY', '')
TZ = ZoneInfo(os.environ.get('TIMEZONE', 'UTC'))

DEFAULT_ROSTER = [
    'Aniket', 'Arshdeep', 'Alfred', 'Arshvir', 'Aryaman', 'Ayush', 'Bijoy',
    'Bikash', 'Bishu', 'Diego', 'Gurkirat', 'Inam', 'Jaskaran', 'Karish',
    'Leon', 'Manav', 'Kipgen', 'Messi', 'Muheet', 'Nikhil', 'Ninthoi',
    'Omang', 'Patrick', 'Shaiza', 'Ricky', 'Rishikanta', 'Sanathoi', 'Shami',
    'Sohel', 'Suhail', 'Suresh', 'Usham', 'Uvais', 'Vishal', 'Lhungdim'
]

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

app = FastAPI()
api_router = APIRouter(prefix="/api")


# ---------- Helpers ----------
def now_utc():
    return datetime.now(timezone.utc)


def today_str():
    return datetime.now(TZ).strftime('%Y-%m-%d')


def parse_day(date_str: Optional[str]) -> Optional[str]:
    """Validate an optional YYYY-MM-DD query parameter. Returns it unchanged, or None if not given."""
    if not date_str:
        return None
    try:
        datetime.strptime(date_str, '%Y-%m-%d')
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date. Use YYYY-MM-DD.")
    return date_str


def normalize_phone(raw: Optional[str]) -> str:
    """Turn a typed phone number into digits-only international format (what WhatsApp links need).
    '' means 'no number'. Raises ValueError with a readable message when it cannot be a valid number."""
    s = (raw or "").strip()
    if not s:
        return ""
    had_plus = s.startswith("+")
    digits = re.sub(r"\D", "", s)
    if not had_plus:
        if digits.startswith("00"):
            digits = digits[2:]                                   # 0091 98765... -> 91 98765...
        elif digits.startswith("0") and len(digits) == 11:
            digits = DEFAULT_COUNTRY_CODE + digits[1:]            # 098765 43210 -> 91 98765 43210
        elif len(digits) == 10:
            digits = DEFAULT_COUNTRY_CODE + digits                # 98765 43210 -> 91 98765 43210
    if digits.startswith("0") or not (8 <= len(digits) <= 15):
        raise ValueError(f"'{s}' is not a valid phone number. Include the country code, e.g. +91 98765 43210.")
    return digits


def norm_name(name: str) -> str:
    return " ".join((name or "").split()).casefold()


_IMPORT_DELIM = re.compile(r"^\s*(.+?)\s*[,;\t:|\-\u2013\u2014]+\s*(\+?\d[\d\s().\-]{6,})\s*$")
_IMPORT_SPACE = re.compile(r"^\s*(.+?)\s+(\+?\d[\d\s().\-]{6,})\s*$")


def parse_import_line(line: str):
    """'Name, 98765 43210' / 'Name - +91 98765 43210' / 'Name<TAB>9876543210' -> (name, number) or None."""
    m = _IMPORT_DELIM.match(line) or _IMPORT_SPACE.match(line)
    return (m.group(1).strip(), m.group(2).strip()) if m else None


FILLS = {
    "green": PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
    "amber": PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
    "red": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
    "grey": PatternFill(start_color="F1F1F1", end_color="F1F1F1", fill_type="solid"),
}


def good_scale_status(v):
    n = v or 0
    if not n:
        return "grey"
    if n <= 2:
        return "red"
    if n == 3:
        return "amber"
    return "green"


def load_status(v):
    n = v or 0
    if not n:
        return "grey"
    if n > 600:
        return "red"
    if n >= 300:
        return "amber"
    return "green"


def soreness_status(areas, sev):
    if not areas:
        return "green"
    n = sev or 0
    if n >= 4:
        return "red"
    if n == 3:
        return "amber"
    return "green"


def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))


def create_token(user_id: str, email: str) -> str:
    payload = {
        "sub": user_id, "email": email,
        "exp": now_utc() + timedelta(days=7), "type": "access",
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(request: Request) -> dict:
    auth_header = request.headers.get("Authorization", "")
    token = auth_header[7:] if auth_header.startswith("Bearer ") else None
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
        if not user:
            raise HTTPException(status_code=401, detail="User not found")
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Session expired, please log in again")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")


# ---------- Models ----------
class LoginRequest(BaseModel):
    email: str
    password: str


class PasswordChangePayload(BaseModel):
    currentPassword: str
    newPassword: str


class AlertPrefsPayload(BaseModel):
    receiveAlerts: bool


class TeamMemberCreate(BaseModel):
    email: str
    password: str
    name: Optional[str] = "Medical Team"


class PhonePayload(BaseModel):
    phone: str = ""


class PhoneImportPayload(BaseModel):
    text: str = ""


class AthleteCreate(BaseModel):
    name: str


class CheckInPayload(BaseModel):
    name: str
    sleepQuality: int
    sleepNotes: Optional[str] = ""
    hydration: int
    motivation: int
    feelingIll: bool = False
    symptoms: List[str] = Field(default_factory=list)
    illnessSeverity: Optional[str] = ""
    temperature: Optional[str] = ""
    illnessNotes: Optional[str] = ""
    sorenessAreas: List[str] = Field(default_factory=list)
    sorenessSide: Optional[Literal["", "Right", "Left", "Both"]] = ""
    sorenessSeverity: Optional[int] = 0
    sorenessNotes: Optional[str] = ""
    logSession: bool = False
    sessionType: Optional[str] = ""
    rpe: Optional[int] = 0
    duration: Optional[int] = 0
    notes: Optional[str] = ""


# ---------- Email ----------
async def alert_recipients() -> List[str]:
    """Everyone who should get alert/report emails: every Medical Team account that hasn't opted out,
    falling back to COACH_EMAIL so a fresh deploy never silently ends up with zero recipients."""
    docs = await db.users.find({"receiveAlerts": {"$ne": False}}, {"_id": 0, "email": 1}).to_list(1000)
    emails = [d["email"] for d in docs if d.get("email")]
    return emails or ([COACH_EMAIL] if COACH_EMAIL else [])


def _send_email_sync(subject: str, html: str, attachments=None, to=None):
    recipients = to if to else ([COACH_EMAIL] if COACH_EMAIL else [])
    if not resend.api_key or not recipients:
        logger.warning("Resend not configured; skipping email.")
        return None
    params = {"from": SENDER_EMAIL, "to": recipients, "subject": subject, "html": html}
    if attachments:
        params["attachments"] = attachments
    return resend.Emails.send(params)


async def send_email(subject: str, html: str, attachments=None, to=None):
    try:
        if to is None:
            to = await alert_recipients()
        return await asyncio.to_thread(_send_email_sync, subject, html, attachments, to)
    except Exception as e:
        logger.error(f"Failed to send email: {e}")
        return None


async def notify_coach_alert(data: CheckInPayload, high_soreness: bool):
    reasons = []
    if data.feelingIll:
        reasons.append("Feeling ill" + (f" ({data.illnessSeverity})" if data.illnessSeverity else ""))
    if high_soreness:
        reasons.append(f"Significant soreness (level {data.sorenessSeverity}/5)")
    rows = "".join(f"<li style='margin:4px 0;'>{r}</li>" for r in reasons)
    html = f"""
    <div style="font-family:Arial,sans-serif;max-width:560px;margin:auto;">
      <div style="background:#0B0F17;color:#fff;padding:20px;border-radius:12px 12px 0 0;">
        <h2 style="margin:0;color:#EF4444;">⚠️ Load and Recovery Monitoring Alert — {data.name}</h2>
      </div>
      <div style="border:1px solid #eee;border-top:none;padding:20px;border-radius:0 0 12px 12px;">
        <p>{data.name} just submitted a check-in flagged for follow-up:</p>
        <ul style="color:#333;">{rows}</ul>
        <table style="width:100%;font-size:14px;color:#333;border-collapse:collapse;">
          <tr><td style="padding:6px 0;color:#888;">Symptoms</td><td>{', '.join(data.symptoms) or '—'}</td></tr>
          <tr><td style="padding:6px 0;color:#888;">Temperature</td><td>{data.temperature or '—'}</td></tr>
          <tr><td style="padding:6px 0;color:#888;">Illness notes</td><td>{data.illnessNotes or '—'}</td></tr>
          <tr><td style="padding:6px 0;color:#888;">Soreness areas</td><td>{', '.join(data.sorenessAreas) or '—'}</td></tr>
          <tr><td style="padding:6px 0;color:#888;">Soreness side</td><td>{data.sorenessSide or '—'}</td></tr>
          <tr><td style="padding:6px 0;color:#888;">Soreness notes</td><td>{data.sorenessNotes or '—'}</td></tr>
        </table>
        <p style="color:#888;font-size:13px;">View the full squad status on your Medical Team Dashboard.</p>
      </div>
    </div>"""
    await send_email(f"⚠️ Load and Recovery Monitoring alert — {data.name}", html)


# ---------- Excel ----------
async def build_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1B4332", end_color="1B4332", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    def style_header(ws):
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

    # Roster
    ws = wb.active
    ws.title = "Roster"
    ws.append(["Athlete Name"])
    roster = await db.roster.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    for a in roster:
        ws.append([a["name"]])
    style_header(ws)

    # Daily
    ws2 = wb.create_sheet("Daily Check-Ins")
    daily_headers = ["Timestamp", "Date", "Athlete", "Sleep (1-5)", "Sleep Notes", "Hydration (1-5)",
                     "Motivation (1-5)", "Feeling Ill", "Symptoms", "Illness Severity", "Temperature (C)",
                     "Illness Notes", "Soreness Areas", "Soreness Side", "Soreness Severity", "Soreness Notes"]
    ws2.append(daily_headers)
    dailies = await db.checkins.find({}, {"_id": 0}).sort("timestamp", 1).to_list(None)
    for d in dailies:
        ws2.append([
            d.get("timestamp", ""), d.get("date", ""), d.get("name", ""),
            d.get("sleepQuality", ""), d.get("sleepNotes", ""), d.get("hydration", ""),
            d.get("motivation", ""), "Yes" if d.get("feelingIll") else "No",
            ", ".join(d.get("symptoms", [])), d.get("illnessSeverity", ""), d.get("temperature", ""),
            d.get("illnessNotes", ""), ", ".join(d.get("sorenessAreas", [])), d.get("sorenessSide", ""),
            d.get("sorenessSeverity", ""), d.get("sorenessNotes", ""),
        ])
    style_header(ws2)

    # RPE
    ws3 = wb.create_sheet("RPE Log")
    rpe_headers = ["Timestamp", "Date", "Athlete", "Session Type", "RPE (1-10)",
                   "Duration (min)", "Notes", "Training Load"]
    ws3.append(rpe_headers)
    sessions = await db.sessions.find({}, {"_id": 0}).sort("timestamp", 1).to_list(None)
    for s in sessions:
        ws3.append([
            s.get("timestamp", ""), s.get("date", ""), s.get("name", ""), s.get("sessionType", ""),
            s.get("rpe", ""), s.get("duration", ""), s.get("notes", ""), s.get("load", ""),
        ])
    style_header(ws3)

    # Dashboard (today, traffic-light colored) — inserted as first sheet
    dash = await compute_dashboard()
    ws4 = wb.create_sheet("Dashboard", 0)
    ws4.append([f"Today's Squad Status — {dash['date']}"])
    ws4["A1"].font = Font(bold=True, size=14)
    ws4.append([])
    dash_headers = ["Athlete", "Sleep", "Hydration", "Motivation", "Illness", "Soreness", "Session Load", "Checked In At", "7-Day Load", "Monotony", "Strain"]
    ws4.append(dash_headers)
    for cell in ws4[3]:
        cell.fill = header_fill
        cell.font = header_font
    for a in dash["athletes"]:
        if not a["checkedIn"]:
            ws4.append([a["name"], "—", "—", "—", "Not checked in", "—", a.get("load") or "—", "—",
                        a.get("weekLoad", 0), a.get("monotony") if a.get("monotony") is not None else "—", a.get("strain") if a.get("strain") is not None else "—"])
            r = ws4.max_row
            for c in range(2, 9):
                ws4.cell(row=r, column=c).fill = FILLS["grey"]
            continue
        ill_text = ("Yes — " + ", ".join(a.get("symptoms", []))) if a.get("feelingIll") else "No"
        soreness_text = ", ".join(a.get("sorenessAreas", [])) if a.get("sorenessAreas") else "None"
        if a.get("sorenessAreas") and a.get("sorenessSide"):
            soreness_text += f" ({a['sorenessSide']})"
        ci = a.get("checkedInAt")
        checked_at = datetime.fromisoformat(ci).astimezone(TZ).strftime("%H:%M") if ci else "—"
        ws4.append([a["name"], a["sleep"], a["hydration"], a["motivation"], ill_text, soreness_text, a.get("load") or "—", checked_at,
                    a.get("weekLoad", 0), a.get("monotony") if a.get("monotony") is not None else "—", a.get("strain") if a.get("strain") is not None else "—"])
        r = ws4.max_row
        ws4.cell(row=r, column=2).fill = FILLS[good_scale_status(a["sleep"])]
        ws4.cell(row=r, column=3).fill = FILLS[good_scale_status(a["hydration"])]
        ws4.cell(row=r, column=4).fill = FILLS[good_scale_status(a["motivation"])]
        ws4.cell(row=r, column=5).fill = FILLS["red" if a.get("feelingIll") else "green"]
        ws4.cell(row=r, column=6).fill = FILLS[soreness_status(a.get("sorenessAreas"), a.get("sorenessSeverity"))]
        ws4.cell(row=r, column=7).fill = FILLS[load_status(a.get("load"))]
    for r in range(4, ws4.max_row + 1):
        risk = next((x.get("monotonyRisk") for x in dash["athletes"] if x["name"] == ws4.cell(row=r, column=1).value), None)
        if risk:
            ws4.cell(row=r, column=10).fill = FILLS[risk]
    ws4.column_dimensions["A"].width = 18
    ws4.column_dimensions["E"].width = 22
    ws4.column_dimensions["F"].width = 24

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- Health ----------
@api_router.get("/health")
async def health():
    return {"status": "ok"}


# ---------- Auth Routes ----------
@api_router.post("/auth/login")
async def login(payload: LoginRequest):
    email = payload.email.strip().lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    token = create_token(user["id"], user["email"])
    return {"access_token": token, "user": {"id": user["id"], "email": user["email"], "name": user.get("name", "Coach")}}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


@api_router.put("/auth/password")
async def change_password(payload: PasswordChangePayload, user: dict = Depends(get_current_user)):
    doc = await db.users.find_one({"id": user["id"]})
    if not doc or not verify_password(payload.currentPassword, doc["password_hash"]):
        raise HTTPException(status_code=400, detail="Current password is incorrect")
    if len(payload.newPassword) < 8:
        raise HTTPException(status_code=400, detail="New password must be at least 8 characters")
    await db.users.update_one({"id": user["id"]}, {"$set": {
        "password_hash": hash_password(payload.newPassword), "passwordManaged": True,
    }})
    return {"status": "ok"}


@api_router.put("/auth/alerts")
async def set_alert_prefs(payload: AlertPrefsPayload, user: dict = Depends(get_current_user)):
    await db.users.update_one({"id": user["id"]}, {"$set": {"receiveAlerts": payload.receiveAlerts}})
    return {"status": "ok", "receiveAlerts": payload.receiveAlerts}


# ---------- Team (Medical Team accounts: multiple logins, each with their own password + alert preference) ----------
@api_router.get("/team")
async def list_team(user: dict = Depends(get_current_user)):
    docs = await db.users.find({}, {"_id": 0, "password_hash": 0}).sort("email", 1).to_list(200)
    for d in docs:
        d.setdefault("receiveAlerts", True)
    return docs


@api_router.post("/team")
async def add_team_member(payload: TeamMemberCreate, user: dict = Depends(get_current_user)):
    email = payload.email.strip().lower()
    if not email or "@" not in email or " " in email:
        raise HTTPException(status_code=400, detail="Enter a valid email address")
    if len(payload.password) < 8:
        raise HTTPException(status_code=400, detail="Password must be at least 8 characters")
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="That email is already on the Medical Team")
    doc = {
        "id": str(uuid.uuid4()), "email": email, "password_hash": hash_password(payload.password),
        "name": (payload.name or "").strip() or "Medical Team", "role": "coach", "receiveAlerts": True,
        "created_at": now_utc().isoformat(),
    }
    await db.users.insert_one(doc)
    return {"id": doc["id"], "email": doc["email"], "name": doc["name"], "receiveAlerts": True}


@api_router.delete("/team/{member_id}")
async def remove_team_member(member_id: str, user: dict = Depends(get_current_user)):
    total = await db.users.count_documents({})
    if total <= 1:
        raise HTTPException(status_code=400, detail="At least one Medical Team account must remain")
    res = await db.users.delete_one({"id": member_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Team member not found")
    return {"status": "ok"}


# ---------- Roster Routes ----------
@api_router.get("/roster")
async def get_roster():
    roster = await db.roster.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    return roster


@api_router.get("/roster/status")
async def roster_status():
    """Public, read-only: which roster athletes have ALREADY checked in today (name + yes/no only,
    nothing about their answers). Used by the check-in form to hide names that have already submitted,
    so a player can't accidentally fill the form twice for the same day."""
    d = today_str()
    roster = await db.roster.find({}, {"_id": 0, "id": 1, "name": 1}).sort("name", 1).to_list(1000)
    checked_today = set(await db.checkins.distinct("name", {"date": d}))
    return [{"id": a["id"], "name": a["name"], "checkedIn": a["name"] in checked_today} for a in roster]


@api_router.post("/roster")
async def add_athlete(payload: AthleteCreate, user: dict = Depends(get_current_user)):
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name required")
    exists = await db.roster.find_one({"name": name})
    if exists:
        raise HTTPException(status_code=400, detail="Athlete already on roster")
    doc = {"id": str(uuid.uuid4()), "name": name, "created_at": now_utc().isoformat()}
    await db.roster.insert_one(doc)
    return {"id": doc["id"], "name": doc["name"]}


@api_router.delete("/roster/{athlete_id}")
async def remove_athlete(athlete_id: str, user: dict = Depends(get_current_user)):
    res = await db.roster.delete_one({"id": athlete_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Athlete not found")
    await db.athlete_contacts.delete_one({"athlete_id": athlete_id})
    return {"status": "ok"}


# ---------- Phone numbers (Medical Team only - kept in their own collection, never in the public roster) ----------
async def contact_map():
    docs = await db.athlete_contacts.find({}, {"_id": 0}).to_list(5000)
    return {d["athlete_id"]: d.get("phone", "") for d in docs}


@api_router.get("/roster/contacts")
async def roster_contacts(user: dict = Depends(get_current_user)):
    roster = await db.roster.find({}, {"_id": 0, "id": 1, "name": 1}).sort("name", 1).to_list(1000)
    contacts = await contact_map()
    return [{"id": a["id"], "name": a["name"], "phone": contacts.get(a["id"], "")} for a in roster]


@api_router.put("/roster/{athlete_id}/phone")
async def set_athlete_phone(athlete_id: str, payload: PhonePayload, user: dict = Depends(get_current_user)):
    athlete = await db.roster.find_one({"id": athlete_id}, {"_id": 0})
    if not athlete:
        raise HTTPException(status_code=404, detail="Athlete not found")
    try:
        phone = normalize_phone(payload.phone)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if phone:
        await db.athlete_contacts.update_one(
            {"athlete_id": athlete_id}, {"$set": {"phone": phone, "updated_at": now_utc().isoformat()}}, upsert=True)
    else:
        await db.athlete_contacts.delete_one({"athlete_id": athlete_id})
    return {"id": athlete_id, "name": athlete["name"], "phone": phone}


@api_router.post("/roster/phones/import")
async def import_phones(payload: PhoneImportPayload, user: dict = Depends(get_current_user)):
    """Paste lines like 'Name, number'. Matches names to the roster (ignoring case) and saves valid numbers."""
    roster = await db.roster.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(1000)
    by_name = {norm_name(a["name"]): a for a in roster}
    updated, unmatched, invalid = [], [], []
    for raw in payload.text.splitlines()[:500]:
        line = raw.strip()
        if not line:
            continue
        parsed = parse_import_line(line)
        if not parsed:
            invalid.append({"line": line, "reason": "Use the format: Name, number"})
            continue
        name, number = parsed
        athlete = by_name.get(norm_name(name))
        if not athlete:
            unmatched.append(line)
            continue
        try:
            phone = normalize_phone(number)
        except ValueError as e:
            invalid.append({"line": line, "reason": str(e)})
            continue
        await db.athlete_contacts.update_one(
            {"athlete_id": athlete["id"]}, {"$set": {"phone": phone, "updated_at": now_utc().isoformat()}}, upsert=True)
        updated.append({"id": athlete["id"], "name": athlete["name"], "phone": phone})
    return {"updated": updated, "unmatched": unmatched, "invalid": invalid}


# ---------- Check-In Route ----------
@api_router.post("/checkin")
async def submit_checkin(data: CheckInPayload):
    ts = now_utc().isoformat()
    d = today_str()
    if await db.checkins.find_one({"name": data.name, "date": d}, {"_id": 1}):
        raise HTTPException(status_code=409, detail=f"{data.name} has already checked in today. "
                             f"If this needs correcting, ask your medical team to edit or delete the entry.")
    checkin_doc = {
        "id": str(uuid.uuid4()), "timestamp": ts, "date": d, "name": data.name,
        "sleepQuality": data.sleepQuality, "sleepNotes": data.sleepNotes,
        "hydration": data.hydration, "motivation": data.motivation,
        "feelingIll": data.feelingIll, "symptoms": data.symptoms,
        "illnessSeverity": data.illnessSeverity, "temperature": data.temperature,
        "illnessNotes": data.illnessNotes, "sorenessAreas": data.sorenessAreas,
        "sorenessSide": data.sorenessSide or "",
        "sorenessSeverity": data.sorenessSeverity, "sorenessNotes": data.sorenessNotes,
    }
    await db.checkins.insert_one(checkin_doc)

    logged_session = False
    if data.logSession and data.rpe and data.duration:
        load = int(data.rpe) * int(data.duration)
        session_doc = {
            "id": str(uuid.uuid4()), "timestamp": ts, "date": d, "name": data.name,
            "sessionType": data.sessionType, "rpe": data.rpe, "duration": data.duration,
            "notes": data.notes, "load": load,
        }
        await db.sessions.insert_one(session_doc)
        logged_session = True

    high_soreness = (data.sorenessSeverity or 0) >= 4
    if data.feelingIll or high_soreness:
        await notify_coach_alert(data, high_soreness)

    return {"status": "ok", "loggedSession": logged_session}


# ---------- Dashboard Route ----------
async def compute_acwr_map(ref_day: Optional[str] = None):
    """Acute:Chronic Workload Ratio per athlete. acute=7-day avg daily load, chronic=28-day avg.
    Calculated as of ref_day (YYYY-MM-DD); defaults to today."""
    today = datetime.strptime(ref_day, '%Y-%m-%d').date() if ref_day else datetime.now(TZ).date()
    start = (today - timedelta(days=27)).strftime('%Y-%m-%d')
    end = today.strftime('%Y-%m-%d')
    acute_cut = (today - timedelta(days=6)).strftime('%Y-%m-%d')
    sessions = await db.sessions.find({"date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(100000)
    agg = {}
    for s in sessions:
        nm = s["name"]
        ld = int(s.get("load", 0))
        agg.setdefault(nm, {"acute": 0, "chronic": 0})
        agg[nm]["chronic"] += ld
        if s["date"] >= acute_cut:
            agg[nm]["acute"] += ld
    result = {}
    for nm, v in agg.items():
        chronic_avg = v["chronic"] / 28
        if chronic_avg <= 0:
            result[nm] = {"acwr": None, "risk": None}
            continue
        ratio = round((v["acute"] / 7) / chronic_avg, 2)
        if ratio > 1.5:
            risk = "red"        # sharp spike — injury risk
        elif ratio >= 1.3 or ratio < 0.8:
            risk = "amber"      # elevated ramp or undertraining
        else:
            risk = "green"      # sweet spot 0.8-1.3
        result[nm] = {"acwr": ratio, "risk": risk}
    return result


# Foster (1998) weekly metrics. The bands are rule-of-thumb markers for a conversation, not a diagnosis:
# monotony >= 2.0 is the value most often quoted as "high"; 1.5-2.0 is shown as "watch".
MONOTONY_WATCH = 1.5
MONOTONY_HIGH = 2.0


def week_stats(daily_loads):
    """Weekly metrics from 7 daily load totals (0 = no session that day), oldest -> newest.
    monotony = mean daily load / standard deviation of the daily loads (sample SD, n-1)
    strain   = weekly load x monotony
    monotony/strain are None when the week has no load, or every day is identical (SD = 0 -> undefined)."""
    total = sum(daily_loads)
    out = {"load": total, "daysLogged": sum(1 for x in daily_loads if x > 0),
           "monotony": None, "strain": None, "monotonyRisk": None}
    if total <= 0 or len(daily_loads) < 2:
        return out
    sd = statistics.stdev(daily_loads)
    if sd == 0:
        return out
    mono_raw = (total / len(daily_loads)) / sd
    mono = round(mono_raw, 2)      # classify on the number people SEE, so 2.00 on screen is never shown as "watch"
    out["monotony"] = mono
    out["strain"] = int(round(total * mono_raw))
    out["monotonyRisk"] = "red" if mono >= MONOTONY_HIGH else ("amber" if mono >= MONOTONY_WATCH else "green")
    return out


def daily_load_series(sessions, end_day, n_days):
    """{athlete name: [total load per day for the n_days ending at end_day, oldest first]}"""
    index = {(end_day - timedelta(days=i)).strftime('%Y-%m-%d'): n_days - 1 - i for i in range(n_days)}
    out = {}
    for s in sessions:
        pos = index.get(s.get("date"))
        if pos is None:
            continue
        out.setdefault(s["name"], [0] * n_days)[pos] += int(s.get("load", 0))
    return out


def change_pct(current, previous):
    return int(round((current - previous) / previous * 100)) if previous and previous > 0 else None


async def compute_weekly_metrics(ref_day: Optional[str] = None):
    """Last-7-days metrics per athlete as of ref_day (default today), plus change vs the 7 days before."""
    end_day = datetime.strptime(ref_day, '%Y-%m-%d').date() if ref_day else datetime.now(TZ).date()
    start = (end_day - timedelta(days=13)).strftime('%Y-%m-%d')
    sessions = await db.sessions.find({"date": {"$gte": start, "$lte": end_day.strftime('%Y-%m-%d')}}, {"_id": 0}).to_list(None)
    result = {}
    for name, loads in daily_load_series(sessions, end_day, 14).items():
        st = week_stats(loads[7:])
        st["prevLoad"] = sum(loads[:7])
        st["changePct"] = change_pct(st["load"], st["prevLoad"])
        result[name] = st
    return result


async def compute_dashboard(day: Optional[str] = None):
    d = day or today_str()
    roster = await db.roster.find({}, {"_id": 0}).sort("name", 1).to_list(1000)

    dailies = await db.checkins.find({"date": d}, {"_id": 0}).sort("timestamp", 1).to_list(10000)
    latest = {}
    for row in dailies:
        latest[row["name"]] = row  # last one wins

    sessions = await db.sessions.find({"date": d}, {"_id": 0}).to_list(10000)
    load_by = {}
    for s in sessions:
        load_by[s["name"]] = load_by.get(s["name"], 0) + int(s.get("load", 0))

    athletes = []
    for a in roster:
        name = a["name"]
        s = latest.get(name)
        entry = {"id": a["id"], "name": name, "checkedIn": bool(s), "load": load_by.get(name, None)}
        if s:
            entry.update({
                "sleep": s["sleepQuality"], "hydration": s["hydration"], "motivation": s["motivation"],
                "feelingIll": s["feelingIll"], "symptoms": s.get("symptoms", []),
                "illnessSeverity": s.get("illnessSeverity", ""),
                "sorenessAreas": s.get("sorenessAreas", []),
                "sorenessSide": s.get("sorenessSide", ""),
                "sorenessSeverity": s.get("sorenessSeverity", 0),
                "checkedInAt": s["timestamp"],
            })
        athletes.append(entry)

    checked_in = sum(1 for a in athletes if a["checkedIn"])
    follow_ups = sum(1 for a in athletes if a["checkedIn"] and (a.get("feelingIll") or (a.get("sorenessSeverity") or 0) >= 4))
    high_load = sum(1 for a in athletes if (a.get("load") or 0) > 600)

    acwr_map = await compute_acwr_map(d)
    for a in athletes:
        m = acwr_map.get(a["name"], {"acwr": None, "risk": None})
        a["acwr"] = m["acwr"]
        a["acwrRisk"] = m["risk"]
    weekly = await compute_weekly_metrics(d)
    for a in athletes:
        w = weekly.get(a["name"])
        a["weekLoad"] = w["load"] if w else 0
        a["weekChangePct"] = w["changePct"] if w else None
        a["monotony"] = w["monotony"] if w else None
        a["strain"] = w["strain"] if w else None
        a["monotonyRisk"] = w["monotonyRisk"] if w else None
        a["weekDaysLogged"] = w["daysLogged"] if w else 0
    load_spikes = sum(1 for a in athletes if a.get("acwrRisk") == "red")

    return {
        "date": d,
        "today": today_str(),
        "summary": {"total": len(athletes), "checkedIn": checked_in, "followUps": follow_ups,
                    "highLoad": high_load, "loadSpikes": load_spikes},
        "athletes": athletes,
    }


@api_router.get("/dashboard")
async def dashboard(date: Optional[str] = None, user: dict = Depends(get_current_user)):
    data = await compute_dashboard(parse_day(date))
    contacts = await contact_map()
    for a in data["athletes"]:
        a["phone"] = contacts.get(a["id"], "")
    return data


async def compute_athlete_trends(name: str, days: int = 7, end: Optional[str] = None) -> dict:
    """Per-athlete history for the last `days` days (7-90) ending on `end` (default today):
    daily load / readiness / soreness / illness, plus weekly load, monotony, strain and change vs the previous week.
    Shared by the coach's /trends route and the athlete-facing /my-stats route."""
    days = max(7, min(int(days), 90))
    end_str = parse_day(end)
    end_day = datetime.strptime(end_str, '%Y-%m-%d').date() if end_str else datetime.now(TZ).date()
    span = days + 7   # one extra week so the oldest week can show its change vs the week before it
    date_strs = [(end_day - timedelta(days=i)).strftime('%Y-%m-%d') for i in range(span - 1, -1, -1)]
    window = {"$gte": date_strs[0], "$lte": date_strs[-1]}
    checkins = await db.checkins.find({"name": name, "date": window}, {"_id": 0}).sort("timestamp", 1).to_list(None)
    latest = {}
    for c in checkins:
        latest[c["date"]] = c
    sessions = await db.sessions.find({"name": name, "date": window}, {"_id": 0}).to_list(None)
    load_by = {}
    for s_ in sessions:
        load_by[s_["date"]] = load_by.get(s_["date"], 0) + int(s_.get("load", 0))
    all_days = []
    for ds in date_strs:
        c = latest.get(ds)
        sev = (c.get("sorenessSeverity") or 0) if c else 0
        areas = (c.get("sorenessAreas") or []) if c else []
        all_days.append({
            "date": ds,
            "load": load_by.get(ds, 0),
            "sleep": c["sleepQuality"] if c else None,
            "hydration": c["hydration"] if c else None,
            "motivation": c["motivation"] if c else None,
            "soreness": sev if areas else None,
            "sorenessAreas": areas,
            "sorenessSide": (c.get("sorenessSide") or "") if c else "",
            "ill": bool(c and c.get("feelingIll")),
            "symptoms": (c.get("symptoms") or []) if c else [],
            "illnessSeverity": (c.get("illnessSeverity") or "") if c else "",
        })
    loads = [d["load"] for d in all_days]
    weeks = []
    for j in range(days // 7 - 1, -1, -1):     # oldest week first
        cur = loads[span - 7 * (j + 1): span - 7 * j]
        prev = loads[span - 7 * (j + 2): span - 7 * (j + 1)]
        st = week_stats(cur)
        st["weekStart"] = all_days[span - 7 * (j + 1)]["date"]
        st["weekEnd"] = all_days[span - 7 * j - 1]["date"]
        st["prevLoad"] = sum(prev)
        st["changePct"] = change_pct(st["load"], st["prevLoad"])
        weeks.append(st)
    return {"name": name, "range": days, "end": date_strs[-1], "days": all_days[-days:], "weeks": weeks}


@api_router.get("/trends")
async def trends(name: str, days: int = 7, end: Optional[str] = None, user: dict = Depends(get_current_user)):
    return await compute_athlete_trends(name, days, end)


# ---------- Team access code (athletes' own "My Stats" page) ----------
class TeamCodePayload(BaseModel):
    code: str


async def get_team_code() -> str:
    doc = await db.settings.find_one({"key": "team_access_code"}, {"_id": 0})
    return (doc or {}).get("code", "")


@api_router.get("/team-access-code")
async def team_access_code(user: dict = Depends(get_current_user)):
    return {"code": await get_team_code()}


@api_router.put("/team-access-code")
async def set_team_access_code(payload: TeamCodePayload, user: dict = Depends(get_current_user)):
    code = payload.code.strip()
    if code and not (4 <= len(code) <= 20):
        raise HTTPException(status_code=400, detail="Code must be 4-20 characters (or blank to turn My Stats off)")
    await db.settings.update_one({"key": "team_access_code"}, {"$set": {"code": code}}, upsert=True)
    return {"code": code}


@api_router.get("/my-stats")
async def my_stats(name: str, code: str, days: int = 7, end: Optional[str] = None):
    """Public but code-gated: lets an athlete see their OWN history (load, readiness, illness, soreness) with
    a single team-wide access code set by the coach - no per-athlete login or email. Never lists other athletes'
    data; the caller must already know their own name (from the public roster) and the shared code."""
    team_code = await get_team_code()
    if not team_code:
        raise HTTPException(status_code=403, detail="My Stats isn't turned on yet. Ask your medical team to set an access code.")
    if code.strip() != team_code:
        raise HTTPException(status_code=403, detail="That access code isn't right. Ask your medical team for the current one.")
    if not await db.roster.find_one({"name": name}, {"_id": 1}):
        raise HTTPException(status_code=404, detail="That name isn't on the roster.")
    return await compute_athlete_trends(name, days, end)


HEATMAP_AREAS = ["Lower Back", "Hamstrings", "Quadriceps", "Calves", "Shoulders", "Gluteus", "Groin", "Upper Back / Neck"]


@api_router.get("/heatmap")
async def heatmap(date: Optional[str] = None, days: int = 1, user: dict = Depends(get_current_user)):
    """Team soreness map: how many athletes reported each body area over the last `days` days (1-30) ending on `date`.
    Uses each athlete's latest check-in per day; per athlete and area the highest severity in the window counts."""
    days = max(1, min(int(days), 30))
    end_str = parse_day(date) or today_str()
    end_day = datetime.strptime(end_str, '%Y-%m-%d').date()
    start_str = (end_day - timedelta(days=days - 1)).strftime('%Y-%m-%d')
    rows = await db.checkins.find({"date": {"$gte": start_str, "$lte": end_str}}, {"_id": 0}).sort("timestamp", 1).to_list(None)
    latest = {}
    for r in rows:
        latest[(r["name"], r["date"])] = r          # last check-in of the day wins
    best = {}
    for (name, day_), r in sorted(latest.items(), key=lambda kv: kv[0][1]):   # oldest day first, so ties go to the newer day
        sev = int(r.get("sorenessSeverity") or 0)
        for area in (r.get("sorenessAreas") or []):
            cur = best.get((name, area))
            if cur is None or sev >= cur["severity"]:
                best[(name, area)] = {"severity": sev, "side": r.get("sorenessSide") or ""}
    areas = {a: {"area": a, "athletes": 0, "avgSeverity": 0, "maxSeverity": 0, "right": 0, "left": 0, "both": 0,
                 "unspecified": 0, "names": []} for a in HEATMAP_AREAS}
    sums = {}
    for (name, area), e in best.items():
        a = areas.setdefault(area, {"area": area, "athletes": 0, "avgSeverity": 0, "maxSeverity": 0, "right": 0,
                                    "left": 0, "both": 0, "unspecified": 0, "names": []})
        a["athletes"] += 1
        sums[area] = sums.get(area, 0) + e["severity"]
        a["maxSeverity"] = max(a["maxSeverity"], e["severity"])
        a[{"Right": "right", "Left": "left", "Both": "both"}.get(e["side"], "unspecified")] += 1
        a["names"].append({"name": name, "severity": e["severity"], "side": e["side"]})
    for area, a in areas.items():
        if a["athletes"]:
            a["avgSeverity"] = round(sums[area] / a["athletes"], 1)
        a["names"].sort(key=lambda n: (-n["severity"], n["name"]))
    ordered = sorted(areas.values(), key=lambda a: (-a["athletes"], -a["maxSeverity"], a["area"]))
    return {"start": start_str, "end": end_str, "days": days,
            "athletesChecked": len({n for (n, _d) in latest}),
            "athletesWithSoreness": len({n for (n, _a) in best}),
            "areas": ordered}


@api_router.get("/checkins/dates")
async def checkin_dates(user: dict = Depends(get_current_user)):
    """Days (YYYY-MM-DD) that have at least one response - used to mark the calendar."""
    days = await db.checkins.distinct("date")
    return {"dates": sorted(d for d in days if d)}


@api_router.get("/checkins")
async def list_checkins(user: dict = Depends(get_current_user), limit: int = 200, date: Optional[str] = None):
    q = {"date": parse_day(date)} if date else {}
    dailies = await db.checkins.find(q, {"_id": 0}).sort("timestamp", -1).to_list(limit)
    sessions = await db.sessions.find(q, {"_id": 0}).to_list(100000)
    smap = {(s["name"], s["timestamp"]): s for s in sessions}
    for c in dailies:
        s = smap.get((c.get("name"), c.get("timestamp")))
        c["load"] = s["load"] if s else None
        c["sessionType"] = s["sessionType"] if s else None
        c["rpe"] = s["rpe"] if s else None
        c["duration"] = s["duration"] if s else None
    return dailies


@api_router.delete("/checkins/today")
async def clear_today(user: dict = Depends(get_current_user)):
    d = today_str()
    c = await db.checkins.delete_many({"date": d})
    s = await db.sessions.delete_many({"date": d})
    return {"status": "ok", "deletedCheckins": c.deleted_count, "deletedSessions": s.deleted_count}


class RestorePayload(BaseModel):
    checkin: dict
    session: Optional[dict] = None


@api_router.post("/checkins/restore")
async def restore_checkin(payload: RestorePayload, user: dict = Depends(get_current_user)):
    c = {k: v for k, v in payload.checkin.items() if k != "_id"}
    if not c.get("id") or not c.get("name"):
        raise HTTPException(status_code=400, detail="Invalid response data")
    c.pop("load", None)
    c.pop("sessionType", None)
    c.pop("rpe", None)
    c.pop("duration", None)
    exists = await db.checkins.find_one({"id": c.get("id")})
    if not exists:
        await db.checkins.insert_one(c)
    if payload.session:
        s = {k: v for k, v in payload.session.items() if k != "_id"}
        s_exists = await db.sessions.find_one({"id": s.get("id")})
        if not s_exists:
            await db.sessions.insert_one(s)
    return {"status": "ok", "restored": c.get("id")}


class CheckInEdit(BaseModel):
    sleepQuality: int
    hydration: int
    motivation: int
    feelingIll: bool = False
    illnessSeverity: Optional[str] = None
    sorenessSeverity: Optional[int] = None
    logSession: bool = False
    sessionType: Optional[str] = None
    rpe: Optional[int] = None
    duration: Optional[int] = None


@api_router.put("/checkins/{checkin_id}")
async def edit_checkin(checkin_id: str, data: CheckInEdit, user: dict = Depends(get_current_user)):
    doc = await db.checkins.find_one({"id": checkin_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Response not found")
    await db.checkins.update_one({"id": checkin_id}, {"$set": {
        "sleepQuality": data.sleepQuality, "hydration": data.hydration,
        "motivation": data.motivation, "feelingIll": data.feelingIll,
        "illnessSeverity": data.illnessSeverity, "sorenessSeverity": data.sorenessSeverity,
    }})
    existing = await db.sessions.find_one({"name": doc["name"], "timestamp": doc["timestamp"]}, {"_id": 0})
    if data.logSession and data.rpe and data.duration:
        load = int(data.rpe) * int(data.duration)
        fields = {"sessionType": data.sessionType, "rpe": data.rpe, "duration": data.duration, "load": load}
        if existing:
            await db.sessions.update_one({"id": existing["id"]}, {"$set": fields})
        else:
            await db.sessions.insert_one({
                "id": str(uuid.uuid4()), "timestamp": doc["timestamp"], "date": doc["date"],
                "name": doc["name"], "notes": "", **fields,
            })
    elif existing:
        await db.sessions.delete_one({"id": existing["id"]})
    return {"status": "ok", "updated": checkin_id}


@api_router.delete("/checkins/{checkin_id}")
async def delete_checkin(checkin_id: str, user: dict = Depends(get_current_user)):
    doc = await db.checkins.find_one({"id": checkin_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Response not found")
    session = await db.sessions.find_one({"name": doc["name"], "timestamp": doc["timestamp"]}, {"_id": 0})
    await db.checkins.delete_one({"id": checkin_id})
    await db.sessions.delete_many({"name": doc["name"], "timestamp": doc["timestamp"]})
    return {"status": "ok", "deleted": checkin_id, "checkin": doc, "session": session}


# ---------- PDF ----------
RL_FILLS = {"green": rl_colors.HexColor("#D4EDDA"), "amber": rl_colors.HexColor("#FFF3CD"),
            "red": rl_colors.HexColor("#F8D7DA"), "grey": rl_colors.HexColor("#F1F1F1")}


async def build_pdf_bytes(day: Optional[str] = None) -> bytes:
    """A one-page-per-~25-athletes PDF snapshot of the squad status for `day` (default today):
    the same numbers and traffic-light colors as the Excel Dashboard sheet, laid out to print or share."""
    dash = await compute_dashboard(day)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                             topMargin=12 * mm, bottomMargin=12 * mm, title="Load and Recovery Monitoring Report")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Heading1"], fontSize=16, spaceAfter=2, textColor=rl_colors.HexColor("#0B0F17"))
    sub_style = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=10, textColor=rl_colors.HexColor("#555555"), alignment=TA_LEFT)

    story = [
        Paragraph("Load and Recovery Monitoring Report", title_style),
        Paragraph(dash["date"], sub_style),
        Spacer(1, 4 * mm),
        Paragraph(
            f"{dash['summary']['checkedIn']} of {dash['summary']['total']} athletes checked in &nbsp;&nbsp;·&nbsp;&nbsp; "
            f"{dash['summary']['followUps']} need follow-up &nbsp;&nbsp;·&nbsp;&nbsp; "
            f"{dash['summary']['highLoad']} high training load &nbsp;&nbsp;·&nbsp;&nbsp; "
            f"{dash['summary']['loadSpikes']} ACWR load spikes", sub_style,
        ),
        Spacer(1, 5 * mm),
    ]

    headers = ["Athlete", "Sleep", "Hyd.", "Mood", "Illness", "Soreness", "Session\nLoad", "7-Day\nLoad", "Monotony", "Strain"]
    header_row = [Paragraph(f"<b>{h}</b>", ParagraphStyle("H", parent=styles["Normal"], fontSize=8, textColor=rl_colors.white)) for h in headers]
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10)
    table_rows = [header_row]
    row_fills = []  # (row_index_in_table, {col: color}) - header excluded, so index 0 = first athlete row
    for a in dash["athletes"]:
        fills = {}
        if not a["checkedIn"]:
            table_rows.append([
                Paragraph(a["name"], cell_style), "—", "—", "—", Paragraph("Not checked in", cell_style), "—",
                str(a.get("load") or "—"), str(a.get("weekLoad", 0)),
                str(a.get("monotony")) if a.get("monotony") is not None else "—",
                str(a.get("strain")) if a.get("strain") is not None else "—",
            ])
            for col in range(10):
                fills[col] = RL_FILLS["grey"]
            row_fills.append(fills)
            continue
        ill_text = ("Yes — " + ", ".join(a.get("symptoms", []))) if a.get("feelingIll") else "No"
        soreness_text = ", ".join(a.get("sorenessAreas", [])) if a.get("sorenessAreas") else "None"
        if a.get("sorenessAreas") and a.get("sorenessSide"):
            soreness_text += f" ({a['sorenessSide']})"
        table_rows.append([
            Paragraph(a["name"], cell_style), str(a["sleep"]), str(a["hydration"]), str(a["motivation"]),
            Paragraph(ill_text, cell_style), Paragraph(soreness_text, cell_style), str(a.get("load") or "—"),
            str(a.get("weekLoad", 0)), str(a.get("monotony")) if a.get("monotony") is not None else "—",
            str(a.get("strain")) if a.get("strain") is not None else "—",
        ])
        fills[1] = RL_FILLS[good_scale_status(a["sleep"])]
        fills[2] = RL_FILLS[good_scale_status(a["hydration"])]
        fills[3] = RL_FILLS[good_scale_status(a["motivation"])]
        fills[4] = RL_FILLS["red"] if a.get("feelingIll") else RL_FILLS["green"]
        fills[5] = RL_FILLS[soreness_status(a.get("sorenessAreas"), a.get("sorenessSeverity"))]
        fills[6] = RL_FILLS[load_status(a.get("load"))]
        if a.get("monotonyRisk"):
            fills[8] = RL_FILLS[a["monotonyRisk"]]
        row_fills.append(fills)

    col_widths = [32 * mm, 12 * mm, 12 * mm, 12 * mm, 34 * mm, 34 * mm, 16 * mm, 16 * mm, 16 * mm, 14 * mm]
    table = Table(table_rows, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), rl_colors.HexColor("#1B4332")),
        ("GRID", (0, 0), (-1, -1), 0.4, rl_colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (3, -1), "CENTER"),
        ("ALIGN", (6, 0), (9, -1), "CENTER"),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i, fills in enumerate(row_fills):
        table_row = i + 1  # +1 for the header row
        for col, color in fills.items():
            style_cmds.append(("BACKGROUND", (col, table_row), (col, table_row), color))
    table.setStyle(TableStyle(style_cmds))
    story.append(table)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(
        "7-Day Load, Monotony and Strain cover the 7 days up to this date. Monotony (average daily load / its "
        "day-to-day variation) of 2.0 or more is high, 1.5-2.0 is worth watching. These are rule-of-thumb "
        "markers, not a diagnosis.", sub_style,
    ))
    doc.build(story)
    return buf.getvalue()


# ---------- Export / Report ----------
@api_router.get("/export/excel")
async def export_excel(user: dict = Depends(get_current_user)):
    data = await build_workbook_bytes()
    filename = f"Load and Recovery Monitoring Report - {today_str()}.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@api_router.get("/export/pdf")
async def export_pdf(date: Optional[str] = None, user: dict = Depends(get_current_user)):
    data = await build_pdf_bytes(parse_day(date))
    d = parse_day(date) or today_str()
    filename = f"Load and Recovery Monitoring Report - {d}.pdf"
    return StreamingResponse(
        io.BytesIO(data), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


async def send_daily_report_email():
    data = await build_workbook_bytes()
    d = today_str()
    dash = await compute_dashboard()
    html = f"""
    <div style="font-family:Arial,sans-serif;max-width:560px;margin:auto;">
      <div style="background:#0B0F17;color:#fff;padding:20px;border-radius:12px 12px 0 0;">
        <h2 style="margin:0;color:#10B981;">Load and Recovery Monitoring Report</h2>
        <p style="margin:4px 0 0;color:#94A3B8;">{d}</p>
      </div>
      <div style="border:1px solid #eee;border-top:none;padding:20px;border-radius:0 0 12px 12px;color:#333;">
        <p><strong>{dash['summary']['checkedIn']}</strong> of <strong>{dash['summary']['total']}</strong> athletes checked in today.</p>
        <p>🔴 {dash['summary']['followUps']} need follow-up &nbsp; ⚡ {dash['summary']['highLoad']} high training load</p>
        <p style="color:#888;font-size:13px;">The full Excel workbook (color-coded Dashboard, Roster, Daily Check-Ins, RPE Log) is attached.</p>
      </div>
    </div>"""
    attachments = [{"filename": f"Load and Recovery Monitoring Report - {d}.xlsx", "content": list(data)}]
    return await send_email(f"Load and Recovery Monitoring Report — {d}", html, attachments)


@api_router.post("/reports/send-now")
async def send_report_now(user: dict = Depends(get_current_user)):
    result = await send_daily_report_email()
    if result is None:
        raise HTTPException(status_code=500, detail="Email could not be sent. Check Resend configuration.")
    return {"status": "ok", "sent_to": await alert_recipients()}


class DailyEmailSettings(BaseModel):
    enabled: bool
    hour: int = 20


@api_router.get("/settings/daily-email")
async def get_daily_email(user: dict = Depends(get_current_user)):
    s = await db.settings.find_one({"key": "daily_email"}, {"_id": 0})
    return {"enabled": bool(s and s.get("enabled")), "hour": (s.get("hour") if s else 20), "recipients": await alert_recipients()}


@api_router.post("/settings/daily-email")
async def set_daily_email(payload: DailyEmailSettings, user: dict = Depends(get_current_user)):
    await db.settings.update_one(
        {"key": "daily_email"},
        {"$set": {"key": "daily_email", "enabled": payload.enabled, "hour": payload.hour}},
        upsert=True,
    )
    return {"enabled": payload.enabled, "hour": payload.hour}


async def scheduled_daily_email():
    s = await db.settings.find_one({"key": "daily_email"})
    if not s or not s.get("enabled"):
        return
    now = datetime.now(TZ)
    if now.hour != s.get("hour", 20):
        return
    if s.get("last_sent") == today_str():
        return
    result = await send_daily_report_email()
    if result is not None:
        await db.settings.update_one({"key": "daily_email"}, {"$set": {"last_sent": today_str()}})
        logger.info("Sent scheduled daily report.")


# ---------- Weekly Digest ----------
async def compute_weekly_summary():
    today = datetime.now(TZ).date()
    start = (today - timedelta(days=6)).strftime('%Y-%m-%d')
    roster = await db.roster.find({}, {"_id": 0}).sort("name", 1).to_list(1000)
    checkins = await db.checkins.find({"date": {"$gte": start}}, {"_id": 0}).to_list(100000)
    sessions = await db.sessions.find({"date": {"$gte": start}}, {"_id": 0}).to_list(100000)
    acwr_map = await compute_acwr_map()
    ci_by = {}
    load_by = {}
    for c in checkins:
        ci_by.setdefault(c["name"], []).append(c)
    for s in sessions:
        load_by[s["name"]] = load_by.get(s["name"], 0) + int(s.get("load", 0))

    def avg(items, key):
        vals = [i.get(key) for i in items if i.get(key)]
        return round(sum(vals) / len(vals), 1) if vals else None

    rows = []
    for a in roster:
        nm = a["name"]
        cs = ci_by.get(nm, [])
        m = acwr_map.get(nm, {"acwr": None, "risk": None})
        rows.append({
            "name": nm, "checkins": len(cs), "sleep": avg(cs, "sleepQuality"),
            "hydration": avg(cs, "hydration"), "motivation": avg(cs, "motivation"),
            "load": load_by.get(nm, 0), "acwr": m["acwr"], "risk": m["risk"],
        })
    return {"start": start, "end": today.strftime('%Y-%m-%d'), "athletes": rows}


async def send_weekly_digest_email():
    summary = await compute_weekly_summary()
    RISK_COLOR = {"red": "#F8D7DA", "amber": "#FFF3CD", "green": "#D4EDDA"}
    body_rows = ""
    spikes = 0
    for r in summary["athletes"]:
        bg = RISK_COLOR.get(r["risk"], "#FFFFFF")
        if r["risk"] == "red":
            spikes += 1
        acwr_txt = f"{r['acwr']}" if r["acwr"] is not None else "—"
        body_rows += (
            f"<tr style='background:{bg};'>"
            f"<td style='padding:6px 8px;'>{r['name']}</td>"
            f"<td style='padding:6px 8px;text-align:center;'>{r['checkins']}</td>"
            f"<td style='padding:6px 8px;text-align:center;'>{r['sleep'] if r['sleep'] is not None else '—'}</td>"
            f"<td style='padding:6px 8px;text-align:center;'>{r['hydration'] if r['hydration'] is not None else '—'}</td>"
            f"<td style='padding:6px 8px;text-align:center;'>{r['motivation'] if r['motivation'] is not None else '—'}</td>"
            f"<td style='padding:6px 8px;text-align:center;'>{r['load']}</td>"
            f"<td style='padding:6px 8px;text-align:center;font-weight:bold;'>{acwr_txt}</td>"
            f"</tr>"
        )
    html = f"""
    <div style="font-family:Arial,sans-serif;max-width:680px;margin:auto;">
      <div style="background:#0B0F17;color:#fff;padding:20px;border-radius:12px 12px 0 0;">
        <h2 style="margin:0;color:#10B981;">Weekly Squad Digest</h2>
        <p style="margin:4px 0 0;color:#94A3B8;">{summary['start']} → {summary['end']}</p>
      </div>
      <div style="border:1px solid #eee;border-top:none;padding:20px;border-radius:0 0 12px 12px;color:#333;">
        <p><strong>{spikes}</strong> athlete(s) showing a sharp load spike (ACWR &gt; 1.5) this week.</p>
        <table style="width:100%;border-collapse:collapse;font-size:13px;">
          <thead>
            <tr style="background:#1B4332;color:#fff;">
              <th style="padding:8px;text-align:left;">Athlete</th>
              <th style="padding:8px;">Check-ins</th>
              <th style="padding:8px;">Avg Sleep</th>
              <th style="padding:8px;">Avg Hydration</th>
              <th style="padding:8px;">Avg Motivation</th>
              <th style="padding:8px;">7d Load</th>
              <th style="padding:8px;">ACWR</th>
            </tr>
          </thead>
          <tbody>{body_rows}</tbody>
        </table>
        <p style="color:#888;font-size:12px;margin-top:14px;">ACWR = acute:chronic workload ratio. Sweet spot 0.8–1.3; &gt;1.5 flags elevated injury risk. Row colors mirror your dashboard.</p>
      </div>
    </div>"""
    data = await build_workbook_bytes()
    attachments = [{"filename": f"Load and Recovery Monitoring Report - {summary['end']}.xlsx", "content": list(data)}]
    return await send_email(f"Weekly Squad Digest — {summary['start']} to {summary['end']}", html, attachments)


class WeeklyDigestSettings(BaseModel):
    enabled: bool
    hour: int = 8


@api_router.get("/settings/weekly-digest")
async def get_weekly_digest(user: dict = Depends(get_current_user)):
    s = await db.settings.find_one({"key": "weekly_digest"}, {"_id": 0})
    return {"enabled": bool(s and s.get("enabled")), "hour": (s.get("hour") if s else 8), "recipients": await alert_recipients()}


@api_router.post("/settings/weekly-digest")
async def set_weekly_digest(payload: WeeklyDigestSettings, user: dict = Depends(get_current_user)):
    await db.settings.update_one(
        {"key": "weekly_digest"},
        {"$set": {"key": "weekly_digest", "enabled": payload.enabled, "hour": payload.hour}},
        upsert=True,
    )
    return {"enabled": payload.enabled, "hour": payload.hour}


@api_router.post("/reports/weekly-now")
async def send_weekly_now(user: dict = Depends(get_current_user)):
    result = await send_weekly_digest_email()
    if result is None:
        raise HTTPException(status_code=500, detail="Email could not be sent. Check Resend configuration.")
    return {"status": "ok", "sent_to": await alert_recipients()}


async def scheduled_weekly_digest():
    s = await db.settings.find_one({"key": "weekly_digest"})
    if not s or not s.get("enabled"):
        return
    now = datetime.now(TZ)
    if now.weekday() != 0:  # Monday
        return
    if now.hour != s.get("hour", 8):
        return
    if s.get("last_sent") == today_str():
        return
    result = await send_weekly_digest_email()
    if result is not None:
        await db.settings.update_one({"key": "weekly_digest"}, {"$set": {"last_sent": today_str()}})
        logger.info("Sent scheduled weekly digest.")


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[o.strip().rstrip('/') for o in os.environ.get('CORS_ORIGINS', '*').split(',') if o.strip()],
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------- Startup ----------
@app.on_event("startup")
async def startup():
    if not os.environ.get('ADMIN_EMAIL', '').strip() or not os.environ.get('ADMIN_PASSWORD', ''):
        raise RuntimeError("ADMIN_EMAIL and ADMIN_PASSWORD must be set (refusing to start with a blank coach login)")
    await db.users.create_index("email", unique=True)
    await db.roster.create_index("name", unique=True)
    await db.athlete_contacts.create_index("athlete_id", unique=True)
    # seed coach
    admin_email = os.environ.get('ADMIN_EMAIL', '').strip().lower()
    admin_password = os.environ.get('ADMIN_PASSWORD', '')
    existing = await db.users.find_one({"email": admin_email})
    if existing is None:
        await db.users.insert_one({
            "id": str(uuid.uuid4()), "email": admin_email,
            "password_hash": hash_password(admin_password), "name": "Coach",
            "role": "coach", "receiveAlerts": True, "created_at": now_utc().isoformat(),
        })
        logger.info("Seeded coach account.")
    elif not verify_password(admin_password, existing["password_hash"]) and not existing.get("passwordManaged"):
        # Only re-syncs from ADMIN_PASSWORD until the account owner sets their own password in the app;
        # after that, "passwordManaged" is set and this account is no longer overwritten on every restart.
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_password)}})
    # seed roster
    count = await db.roster.count_documents({})
    if count == 0:
        docs = [{"id": str(uuid.uuid4()), "name": n, "created_at": now_utc().isoformat()} for n in DEFAULT_ROSTER]
        await db.roster.insert_many(docs)
        logger.info("Seeded default roster.")
    # start scheduler for automated daily email (fires hourly, gated by settings)
    scheduler = AsyncIOScheduler(timezone=str(TZ))
    scheduler.add_job(scheduled_daily_email, "cron", minute=0, id="daily_email", replace_existing=True)
    scheduler.add_job(scheduled_weekly_digest, "cron", minute=0, id="weekly_digest", replace_existing=True)
    scheduler.start()
    logger.info("Scheduler started.")


@app.on_event("shutdown")
async def shutdown():
    client.close()
