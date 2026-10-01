"""
GPS report import.

Turns a team GPS report into per-athlete rows and matches those rows to the
team roster by name. Pure functions (bytes/text in, dicts out) so they can be
tested without a database.

Supported inputs
  * PDF  -- Catapult-style "Activity Report" with an "Athlete Breakdown" table
            (Name, Position, Accel&Decel Efforts, Distance, Overall %, High Speed
            Distance, Sprint Efforts, High Metabolic Load Distance, Maximum
            Velocity, Duration, Meterage Per Minute).
  * CSV / XLSX -- any export whose column headers can be recognised (see
            HEADER_ALIASES); unrecognised columns are reported back, not guessed.

Names in these reports wrap and hyphenate across lines ("Bishu, Shar-" / "ma"),
so the PDF path rejoins hyphenated fragments before anything is matched.
"""
import csv
import difflib
import io
import re
import unicodedata
from datetime import datetime
from typing import Dict, List, Optional


class ParseError(Exception):
    """A problem with the uploaded file that the user can act on."""


METRIC_KEYS = (
    "accel_decel", "distance_m", "overall_pct", "hsr_distance_m", "sprints",
    "hml_distance_m", "max_speed_kmh", "duration_min", "m_per_min",
)

# One athlete's metrics, in the order this report prints them:
# position, accel+decel, distance, overall %, HS distance, sprints, HML distance,
# max velocity, duration (h:mm:ss), metres per minute.
_NUM = r"\d+(?:\.\d+)?"
ROW_RE = re.compile(
    r"(?<![A-Za-z])(?P<pos>[A-Za-z]{1,4})\s+"
    rf"(?P<accel>{_NUM})\s+(?P<dist>{_NUM})\s+(?P<overall>{_NUM})\s+"
    rf"(?P<hsd>{_NUM})\s+(?P<sprints>{_NUM})\s+(?P<hml>{_NUM})\s+"
    rf"(?P<vmax>{_NUM})\s+(?P<dur>\d+:\d{{2}}:\d{{2}})\s+(?P<mpm>{_NUM})"
)


# ------------------------------------------------------------------ names
def clean_name(raw: str) -> str:
    """Rejoin hyphen line-wraps and tidy whitespace, keeping the report's own
    word order ("Bishu, Sharma"). A hyphen only counts as a wrap when the
    next fragment starts lowercase, so a real hyphenated name ("Al-Hassan")
    is left alone."""
    s = re.sub(r"-\s*\n\s*(?=[a-z])", "", raw or "")
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"\s*,\s*", ", ", s)
    return s.strip(" ,")


def display_name(cleaned: str) -> str:
    """Name to use when creating a roster player: same order, comma dropped."""
    return re.sub(r"\s+", " ", cleaned.replace(",", " ")).strip()


def name_tokens(name: str) -> List[str]:
    s = unicodedata.normalize("NFKD", name or "")
    s = "".join(ch for ch in s if not unicodedata.combining(ch)).lower()
    return [t for t in re.sub(r"[^a-z0-9]+", " ", s).split() if t]


def match_key(name: str) -> str:
    """Order-insensitive key: "Sharma, Bishu" and "Bishu Sharma" collide."""
    return " ".join(sorted(name_tokens(name)))


# ------------------------------------------------------------------ helpers
def parse_duration_minutes(value) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    s = str(value).strip()
    m = re.fullmatch(r"(\d+):(\d{2}):(\d{2})", s)
    if m:
        h, mi, se = map(int, m.groups())
        return round(h * 60 + mi + se / 60, 2)
    m = re.fullmatch(r"(\d+):(\d{2})", s)  # mm:ss
    if m:
        mi, se = map(int, m.groups())
        return round(mi + se / 60, 2)
    try:
        return round(float(s.replace(",", "")), 2)
    except ValueError:
        return None


def _num(value) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError:
        return None


def _row_from_match(m: "re.Match", raw_name: str) -> dict:
    name = clean_name(raw_name)
    return {
        "raw_name": name, "name": display_name(name), "position": m.group("pos").upper(),
        "accel_decel": float(m.group("accel")), "distance_m": float(m.group("dist")),
        "overall_pct": float(m.group("overall")), "hsr_distance_m": float(m.group("hsd")),
        "sprints": float(m.group("sprints")), "hml_distance_m": float(m.group("hml")),
        "max_speed_kmh": float(m.group("vmax")), "duration_min": parse_duration_minutes(m.group("dur")),
        "m_per_min": float(m.group("mpm")),
    }


# ------------------------------------------------------------------ PDF
def parse_session_header(text: str) -> dict:
    out = {"name": None, "date": None, "md_tag": None, "total_time": None, "team": None, "venue": None}
    m = re.search(r"[A-Za-z]+DAY,?\s+([A-Za-z]+)\s+(\d{1,2}),\s*(\d{4})", text, re.I)
    if m:
        try:
            out["date"] = datetime.strptime(f"{m.group(1)} {m.group(2)} {m.group(3)}", "%B %d %Y").strftime("%Y-%m-%d")
        except ValueError:
            pass
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    for line in lines[1:4]:
        if re.search(r"TOTAL TIME|^Summary", line, re.I):
            continue
        out["name"] = line
        break
    if not out["date"] and out["name"]:
        d = re.search(r"(\d{2})-(\d{2})-(\d{4})", out["name"])
        if d:
            out["date"] = f"{d.group(3)}-{d.group(2)}-{d.group(1)}"
    m = re.search(r"\b(MD\s*[-+]?\s*\d+)\b", text)
    if m:
        out["md_tag"] = re.sub(r"\s+", "", m.group(1))
    m = re.search(r"TOTAL TIME\s+(\d+:\d{2}:\d{2})", text)
    if m:
        out["total_time"] = m.group(1)
    m = re.search(r"TEAM\s+(.+?)\s+VENUE", text)
    if m:
        out["team"] = m.group(1).strip()
    m = re.search(r"VENUE\s+(.+?)\s*$", text, re.M)
    if m:
        out["venue"] = m.group(1).strip()
    return out


def _rows_from_table_cells(tables) -> List[dict]:
    rows = []
    for table in tables:
        for cells in table:
            text = "\n".join(str(c) for c in cells if c)
            if re.match(r"\s*Average", text, re.I):
                continue
            m = ROW_RE.search(text)
            if not m:
                continue  # header rows, blank rows
            rows.append(_row_from_match(m, text[: m.start()].strip() + "\n" + text[m.end():].strip()))
    return rows


def _rows_from_text(text: str) -> List[dict]:
    """Fallback for reports whose tables have no ruling lines.

    A wrapped name has its first line ABOVE the metrics line and the rest
    BELOW it. A row's name is incomplete if its first part ends in ',' or '-';
    only then does the next stray line belong to it rather than to the next row.
    """
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    start = 0
    for i, l in enumerate(lines[:14]):
        if re.search(r"\(km/h\)|Minute$", l):
            start = i + 1
    rows: List[dict] = []
    firsts: List[str] = []
    prev = None  # (match, before, after) of the previous row
    parsed: List[list] = []
    for line in lines[start:]:
        if re.match(r"Average", line, re.I):
            continue
        m = ROW_RE.search(line)
        if not m:
            firsts.append(line)
            continue
        if parsed and firsts and re.search(r"[,-]$", parsed[-1][1]):
            parsed[-1][2] = firsts.pop(0)
        prefix = line[: m.start()].strip()
        before = " ".join(firsts + ([prefix] if prefix else []))
        parsed.append([m, before, line[m.end():].strip()])
        firsts = []
    if parsed and firsts and re.search(r"[,-]$", parsed[-1][1]):
        parsed[-1][2] = firsts[0]
    for m, before, after in parsed:
        rows.append(_row_from_match(m, before + "\n" + after))
    return rows


def parse_pdf(data: bytes) -> dict:
    try:
        import pdfplumber
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception:
        raise ParseError("That file couldn't be opened as a PDF.")
    with pdf:
        texts = [(p.extract_text() or "") for p in pdf.pages]
        session = parse_session_header(texts[0] if texts else "")
        athletes: List[dict] = []
        for page, text in zip(pdf.pages, texts):
            if not re.search(r"Athlete\s+Breakdown", text, re.I):
                continue
            page_rows = _rows_from_table_cells(page.extract_tables())
            if not page_rows:
                page_rows = _rows_from_text(text)
            athletes.extend(page_rows)
    if not athletes:
        raise ParseError(
            "Couldn't find an 'Athlete Breakdown' table in this PDF. It needs the per-athlete table "
            "(Name, Position, Distance, …). If your system exports a different layout, upload it as CSV or Excel."
        )
    return {"format": "pdf", "session": session, "athletes": athletes, "unmapped_columns": []}


# ------------------------------------------------------------------ CSV / XLSX
HEADER_ALIASES = {
    "name": ["name", "athlete", "athlete name", "player", "player name"],
    "position": ["position", "pos", "role"],
    "accel_decel": ["accel&decel efforts", "accel decel efforts", "accel+decel", "accel+decel #", "accelerations + decelerations",
                    "acc+dec", "accel/decel", "accel decel", "accel_decel", "accel+decel efforts"],
    "distance_m": ["distance (m)", "distance", "total distance", "total distance (m)", "dist (m)", "distance m"],
    "overall_pct": ["overall (%)", "overall"],
    "hsr_distance_m": ["high speed distance (m)", "high speed distance", "hsr", "hsr distance", "hsr distance (m)",
                       "hs distance", "hs dist (m)", "high speed running", "high speed running (m)"],
    "sprints": ["sprint efforts", "sprints", "sprint count", "sprint"],
    "hml_distance_m": ["high metabolic load distance (m)", "high metabolic load distance", "hml distance", "hml distance (m)", "hmld"],
    "max_speed_kmh": ["maximum velocity (km/h)", "maximum velocity", "max velocity", "max speed", "max speed (km/h)", "top speed"],
    "duration": ["duration", "time", "total time", "duration (min)", "minutes"],
    "m_per_min": ["meterage per minute", "metrage per minute", "m/min", "meters per minute", "metres per minute",
                  "meterage per min", "relative distance", "distance per minute"],
    "date": ["date", "session date"],
    "session": ["session", "session name", "activity"],
}
_ALIAS_LOOKUP = {}
for _key, _names in HEADER_ALIASES.items():
    for _n in _names:
        _ALIAS_LOOKUP[re.sub(r"[^a-z0-9%/+&]+", " ", _n.lower()).strip()] = _key


def _norm_header(h) -> str:
    return re.sub(r"[^a-z0-9%/+&]+", " ", str(h or "").lower()).strip()


def _parse_date_cell(value) -> Optional[str]:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    s = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _rows_from_table(table: List[list]) -> dict:
    header_idx, col_map = None, {}
    for i, row in enumerate(table[:15]):
        cm = {}
        for j, cell in enumerate(row):
            key = _ALIAS_LOOKUP.get(_norm_header(cell))
            if key and key not in cm.values():
                cm[j] = key
        if "name" in cm.values() and len(cm) >= 3:
            header_idx, col_map = i, cm
            break
    if header_idx is None:
        raise ParseError(
            "Couldn't recognise the column headers. The file needs at least a Name column plus a few metrics "
            "such as Distance, High Speed Distance, Sprint Efforts."
        )
    unmapped = [str(c) for j, c in enumerate(table[header_idx]) if j not in col_map and str(c or "").strip()]
    athletes = []
    for row in table[header_idx + 1:]:
        rec = {col_map[j]: row[j] for j in col_map if j < len(row)}
        name = str(rec.get("name") or "").strip()
        if not name or re.match(r"average|total|mean", name, re.I):
            continue
        item = {"raw_name": clean_name(name), "name": display_name(clean_name(name)),
                "position": str(rec.get("position") or "").strip().upper()}
        for k in METRIC_KEYS:
            if k == "duration_min":
                item[k] = parse_duration_minutes(rec.get("duration"))
            else:
                item[k] = _num(rec.get(k))
        if rec.get("date"):
            item["date"] = _parse_date_cell(rec["date"])
        if rec.get("session"):
            item["session"] = str(rec["session"]).strip()
        athletes.append(item)
    if not athletes:
        raise ParseError("The file has headers but no athlete rows.")
    return {"athletes": athletes, "unmapped_columns": unmapped}


def parse_csv(data: bytes) -> dict:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    result = _rows_from_table(list(csv.reader(io.StringIO(text), dialect)))
    return {"format": "csv", "session": {"name": None, "date": None, "md_tag": None, "total_time": None}, **result}


def parse_xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ParseError("That file couldn't be opened as an Excel workbook (.xlsx).")
    table = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    result = _rows_from_table(table)
    return {"format": "xlsx", "session": {"name": None, "date": None, "md_tag": None, "total_time": None}, **result}


def parse_file(filename: str, data: bytes) -> dict:
    name = (filename or "").lower()
    if data[:4] == b"%PDF" or name.endswith(".pdf"):
        parsed = parse_pdf(data)
    elif name.endswith(".xlsx"):
        parsed = parse_xlsx(data)
    elif name.endswith(".csv") or name.endswith(".txt"):
        parsed = parse_csv(data)
    else:
        raise ParseError("Upload a PDF report, or a CSV / Excel (.xlsx) export.")
    if not parsed["session"].get("date"):
        dates = {a.get("date") for a in parsed["athletes"] if a.get("date")}
        if len(dates) == 1:
            parsed["session"]["date"] = dates.pop()
    return parsed


# ------------------------------------------------------------------ matching
AUTO_THRESHOLD = 0.75      # below this a match is only ever a suggestion
SUGGEST_THRESHOLD = 0.5
TIE_MARGIN = 0.04          # a match is ambiguous if a rival is this close


def _token_ratio(t: str, pool) -> float:
    return max((difflib.SequenceMatcher(None, t, p).ratio() for p in pool), default=0.0)


def score_names(a_tokens: List[str], b_tokens: List[str]) -> float:
    A, B = set(a_tokens), set(b_tokens)
    if not A or not B:
        return 0.0
    if A == B:
        return 1.0
    small, large = (A, B) if len(A) <= len(B) else (B, A)
    if all(len(t) >= 3 for t in small) and small <= large:
        return 0.86 if len(small) == 1 else 0.88
    ratios = [_token_ratio(t, large) for t in small]
    if all(r >= 0.86 for r in ratios) and all(len(t) >= 3 for t in small):
        return 0.78 if len(small) == len(large) else 0.75
    whole = difflib.SequenceMatcher(None, " ".join(sorted(A)), " ".join(sorted(B))).ratio()
    return round(whole * 0.8, 3) if whole >= 0.8 else round(whole * 0.6, 3)


def match_athletes(rows: List[dict], roster: List[dict]) -> List[dict]:
    """roster: [{id, name, aliases: [match_key, ...]}].

    Returns one dict per row: {"match": {...}|None, "suggestions": [...]}.
    A row is only auto-matched when it is BOTH that player's clear best fit and
    the row that player fits best -- otherwise (two rows fighting over one
    player, or one row fitting two players) it is left for the coach to decide,
    because a silent wrong match would put one athlete's data on another.
    """
    n, m = len(rows), len(roster)
    S = [[0.0] * m for _ in range(n)]
    how = [[""] * m for _ in range(n)]
    for i, row in enumerate(rows):
        key = match_key(row["raw_name"])
        toks = name_tokens(row["raw_name"])
        for j, p in enumerate(roster):
            if key and key in (p.get("aliases") or []):
                S[i][j], how[i][j] = 1.0, "alias"
                continue
            s = score_names(toks, name_tokens(p["name"]))
            S[i][j] = s
            how[i][j] = "exact" if s == 1.0 else ("partial" if s >= 0.86 else "fuzzy")

    def runner_up(vals, best_idx):
        others = [v for k, v in enumerate(vals) if k != best_idx]
        return max(others) if others else 0.0

    results = []
    for i in range(n):
        order = sorted(range(m), key=lambda j: -S[i][j])
        suggestions = [{"player_id": roster[j]["id"], "name": roster[j]["name"], "score": round(S[i][j], 2)}
                       for j in order[:3] if S[i][j] >= SUGGEST_THRESHOLD]
        match = None
        if order:
            j = order[0]
            col = [S[r][j] for r in range(n)]
            best = S[i][j]
            clear_for_row = best - runner_up(S[i], j) >= TIE_MARGIN
            clear_for_player = best >= max(col) and best - runner_up(col, i) >= TIE_MARGIN
            if best >= AUTO_THRESHOLD and clear_for_row and clear_for_player:
                match = {"player_id": roster[j]["id"], "name": roster[j]["name"], "score": round(best, 2), "how": how[i][j]}
        results.append({"match": match, "suggestions": suggestions})
    return results
